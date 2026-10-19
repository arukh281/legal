"""Case Import Service: Excel/CSV parsing, normalization, preview, and atomic commit.

Normative sources:
- Session R3 specification & review approvals:
  1. Openpyxl in strict read-only, data-only mode; reject .xlsm and VBA macros.
  2. 10 MB file limit; 100 MB uncompressed limit; 5,000 rows limit.
  3. Day-first Indian date parsing with ambiguity warnings and serial number support.
  4. Deduplication: CIN/PAN -> normalized name; readable/unreadable case numbers and title matching.
  5. Importer + lead lawyer email both become LEAD members.
  6. Atomic commit with pg_advisory_xact_lock per tenant.
  7. CSV report with formula sanitization (=, +, -, @, \\t, \\r).
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
import openpyxl.styles
from django.db import connection, transaction
from django.utils import timezone
from openpyxl.utils.datetime import from_excel

from anchor_lib.case_numbers import normalize_case_number
from workspace.audit import record_audit_event
from workspace.models import AppUser, Client, Hearing, Matter, MatterCaseLink
from workspace.services import (
    create_case_link_service,
    create_client_service,
    create_hearing_service,
    create_matter_service,
)

IST = ZoneInfo("Asia/Kolkata")
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_UNCOMPRESSED_SIZE = 100 * 1024 * 1024  # 100 MB
MAX_ROWS = 5000

TEMPLATE_COLUMNS = [
    "client name",
    "client type",
    "CIN or PAN",
    "client role",
    "matter title",
    "case type",
    "court or forum",
    "case number",
    "next hearing date",
    "lead lawyer email",
    "notes",
]

HEADER_SYNONYMS: dict[str, list[str]] = {
    "client_name": ["client name", "client", "party", "party name", "client_name"],
    "client_type": ["client type", "type", "client_type", "entity type"],
    "cin_or_pan": ["cin or pan", "cin/pan", "cin", "pan", "identifier", "cin_or_pan"],
    "client_role": ["client role", "client_role", "role", "capacity", "client side", "side"],
    "matter_title": ["matter title", "matter", "title", "case name", "matter_name", "matter_title"],
    "case_type": ["case type", "case_type", "nature of case", "type of case"],
    "court_or_forum": ["court or forum", "court", "forum", "bench", "tribunal", "court_or_forum"],
    "case_number": ["case number", "case no", "case nos", "case_number", "cp no", "petition no"],
    "next_hearing_date": [
        "next hearing date",
        "hearing date",
        "next date",
        "date",
        "next_hearing_date",
        "listing date",
    ],
    "lead_lawyer_email": [
        "lead lawyer email",
        "lead lawyer",
        "lawyer email",
        "lead email",
        "lead_lawyer_email",
        "assignee email",
    ],
    "notes": ["notes", "remarks", "comment", "comments", "description"],
}


def normalize_client_name(name: str) -> str:
    """Fold case, punctuation, repeated spaces, leading M/s, and Pvt/Private, Ltd/Limited."""
    if not name:
        return ""
    s = name.strip()
    # Leading M/s, M/S, M/s., Messrs, etc.
    s = re.sub(r"^(?:m\s*/\s*s\.?|messrs\.?)\s+", "", s, flags=re.IGNORECASE)
    # Fold punctuation (dots, commas, hyphens, quotes, slashes, brackets) to spaces
    s = re.sub(r"[.,\-–—'\"/()\\]+", " ", s)
    # Fold Pvt / Private
    s = re.sub(r"\bpvt\b", "private", s, flags=re.IGNORECASE)
    # Fold Ltd / Limited
    s = re.sub(r"\bltd\b", "limited", s, flags=re.IGNORECASE)
    # Fold repeated whitespace
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s


def sanitize_csv_cell(value: Any) -> str:
    """Sanitize cell value for CSV export to prevent formula execution."""
    if value is None:
        return ""
    s = str(value)
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{s}"
    return s


def inspect_file_safety(file_bytes: bytes, filename: str) -> tuple[str, list[str], list[str]]:
    """Inspect file magic bytes, zip safety, and extract sheet names if Excel."""
    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError(f"File size ({len(file_bytes)} bytes) exceeds 10 MB limit.")

    # 1. Check for Zip / Excel (.xlsx)
    if file_bytes.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                # Check uncompressed size
                total_uncompressed = sum(info.file_size for info in z.infolist())
                if total_uncompressed > MAX_UNCOMPRESSED_SIZE:
                    raise ValueError(
                        f"Total uncompressed size ({total_uncompressed} bytes) exceeds 100 MB limit."
                    )

                # Check macros and VBA projects
                namelist = z.namelist()
                for name in namelist:
                    lower_entry = name.lower()
                    if "vbaproject.bin" in lower_entry or "macroenabled" in lower_entry:
                        raise ValueError(
                            "Excel files with macros (.xlsm) or VBA code are strictly forbidden."
                        )

                # Check [Content_Types].xml for macroEnabled declarations
                if "[Content_Types].xml" in namelist:
                    content_types = z.read("[Content_Types].xml").decode("utf-8", errors="ignore")
                    if "macroEnabled" in content_types:
                        raise ValueError(
                            "Excel files with macros (.xlsm) or VBA code are strictly forbidden."
                        )
        except zipfile.BadZipFile as err:
            raise ValueError("Corrupt Excel file archive.") from err

        # Read sheet names using openpyxl in read-only mode
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        wb.close()
        return ("xlsx", sheet_names, [])

    # 2. Check for CSV (plain text)
    warnings: list[str] = []
    # Reject binary null bytes
    if b"\x00" in file_bytes:
        raise ValueError("Binary or corrupted file detected; only .xlsx and .csv are supported.")

    try:
        # UTF-8 with or without BOM
        file_bytes.decode("utf-8-sig")
        return ("csv", ["Sheet1"], warnings)
    except UnicodeDecodeError:
        try:
            file_bytes.decode("cp1252")
            warnings.append("File decoded using Windows-1252 encoding.")
            return ("csv", ["Sheet1"], warnings)
        except UnicodeDecodeError as err:
            raise ValueError(
                "Unsupported file format or encoding; please upload a valid UTF-8 CSV or .xlsx file."
            ) from err


def parse_raw_rows(
    file_bytes: bytes, file_type: str, sheet_name: str | None = None
) -> tuple[list[str], list[list[Any]], list[str]]:
    """Extract headers and raw rows from Excel or CSV."""
    warnings: list[str] = []

    if file_type == "xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets[0]

        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            wb.close()
            return [], [], ["The uploaded worksheet is empty."]

        headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
        raw_rows: list[list[Any]] = []
        for r in rows_iter:
            if not any(cell is not None and str(cell).strip() != "" for cell in r):
                continue  # skip entirely blank lines
            raw_rows.append(list(r))
            if len(raw_rows) > MAX_ROWS:
                wb.close()
                raise ValueError(f"Spreadsheet exceeds the maximum limit of {MAX_ROWS} rows.")
        wb.close()
        return headers, raw_rows, warnings

    elif file_type == "csv":
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = file_bytes.decode("cp1252")
            warnings.append("File decoded using Windows-1252 encoding.")

        csv_reader = csv.reader(io.StringIO(text))
        csv_header_row = next(csv_reader, None)
        if not csv_header_row:
            return [], [], ["The uploaded CSV file is empty."]

        headers = [c.strip() for c in csv_header_row]
        csv_rows: list[list[Any]] = []
        for csv_row in csv_reader:
            if not any(c.strip() != "" for c in csv_row):
                continue
            csv_rows.append(list(csv_row))
            if len(csv_rows) > MAX_ROWS:
                raise ValueError(f"CSV exceeds the maximum limit of {MAX_ROWS} rows.")

        return headers, csv_rows, warnings

    raise ValueError(f"Unsupported file type: {file_type}")


def auto_map_headers(headers: list[str]) -> dict[str, int]:
    """Auto-match uploaded header names to standard fields."""
    mapping: dict[str, int] = {}
    normalized_headers = [h.strip().lower() for h in headers]

    for field, synonyms in HEADER_SYNONYMS.items():
        matched_idx = -1
        for syn in synonyms:
            if syn in normalized_headers:
                matched_idx = normalized_headers.index(syn)
                break
        if matched_idx != -1:
            mapping[field] = matched_idx

    return mapping


def parse_import_date(raw_val: Any, today_d: date) -> tuple[date | None, list[str]]:
    """Parse date respecting Indian day-first practice, flagging ambiguity and past dates."""
    warnings: list[str] = []
    if raw_val is None:
        return None, []

    # 1. Native date/datetime from openpyxl
    if isinstance(raw_val, datetime):
        d = raw_val.date()
        if d < today_d:
            warnings.append("Hearing date is in the past (historical hearing)")
        return d, warnings

    if isinstance(raw_val, date):
        if raw_val < today_d:
            warnings.append("Hearing date is in the past (historical hearing)")
        return raw_val, warnings

    # 2. Excel numeric serial number
    if isinstance(raw_val, (int, float)):
        # Session R3 Fix Pass 2 - Item 7: Treat as serial only between 36526 and 73050 (1 Jan 2000 to 31 Dec 2099)
        if not (36526 <= raw_val <= 73050):
            return None, [f"Unreadable date: '{raw_val}'"]
        try:
            dt = from_excel(raw_val)
            d = dt.date()
            if d < today_d:
                warnings.append("Hearing date is in the past (historical hearing)")
            return d, warnings
        except Exception:
            return None, [f"Unreadable date: '{raw_val}'"]

    # 3. String representation
    s = str(raw_val).strip()
    if not s or s.lower() in ("na", "n/a", "to be listed", "tbd", "-", "none", "null", ""):
        return None, [f"No date provided: '{raw_val}'"]

    # Excel serial number passed as numeric string
    if s.isdigit() and len(s) in (4, 5):
        val_num = int(s)
        if 36526 <= val_num <= 73050:
            try:
                dt = from_excel(val_num)
                d = dt.date()
                if d < today_d:
                    warnings.append("Hearing date is in the past (historical hearing)")
                return d, warnings
            except Exception:
                pass
        elif len(s) == 4:
            # 4-digit numbers like '2024' are not serials; avoid becoming 1905 dates
            return None, [f"Unreadable date: '{raw_val}'"]

    # Try ISO YYYY-MM-DD
    iso_match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$", s)
    if iso_match:
        y, m, day = int(iso_match.group(1)), int(iso_match.group(2)), int(iso_match.group(3))
        try:
            d = date(y, m, day)
            if d < today_d:
                warnings.append("Hearing date is in the past (historical hearing)")
            return d, warnings
        except ValueError:
            return None, [f"Invalid date: '{raw_val}'"]

    # Try Indian Day-First: DD/MM/YYYY, DD.MM.YYYY, DD-MM-YYYY
    day_first_match = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", s)
    if day_first_match:
        p1, p2, p3 = (
            int(day_first_match.group(1)),
            int(day_first_match.group(2)),
            int(day_first_match.group(3)),
        )
        if p3 < 100:
            p3 += 2000

        try:
            d = date(p3, p2, p1)
            # Check ambiguity: both day and month <= 12
            if p1 <= 12 and p2 <= 12:
                weekday_name = d.strftime("%A")
                month_name = d.strftime("%B")
                warnings.append(
                    f"Ambiguous date: interpreted as {weekday_name}, {p1} {month_name} {p3} (day-first)"
                )
            if d < today_d:
                warnings.append("Hearing date is in the past (historical hearing)")
            return d, warnings
        except ValueError:
            return None, [f"Invalid date: '{raw_val}'"]

    # Try alphanumeric month formats: 21-Oct-26, 21-Oct-2026, 21 Oct 2026
    for fmt in ("%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %b %y", "%d.%b.%Y"):
        try:
            dt = datetime.strptime(s, fmt)
            d = dt.date()
            if d < today_d:
                warnings.append("Hearing date is in the past (historical hearing)")
            return d, warnings
        except ValueError:
            continue

    return None, [f"Unreadable date: '{raw_val}'"]


def normalize_client_role(raw_val: Any) -> tuple[str, str | None]:
    """Normalize client role into standard canonical role with warning if unreadable. (Session R3 Fix Pass 2 - Item 3)"""
    if raw_val is None:
        return "NOT_SET", "Client role not specified; set to 'Not set'."
    s = str(raw_val).strip()
    if not s:
        return "NOT_SET", "Client role not specified; set to 'Not set'."

    # Ignore suffixes like 'No. 2', 'No 1', etc.
    cleaned = re.sub(r"\s+no\.?\s*\d+.*$", "", s, flags=re.IGNORECASE).strip()
    cleaned_lower = cleaned.lower()

    ROLE_MAP = {
        "petitioner": "PETITIONER",
        "applicant": "PETITIONER",
        "plaintiff": "PETITIONER",
        "complainant": "PETITIONER",
        "respondent": "RESPONDENT",
        "defendant": "RESPONDENT",
        "appellant": "APPELLANT",
        "claimant": "CLAIMANT",
        "advisory": "ADVISORY",
        "corporate": "ADVISORY",
        "not_set": "NOT_SET",
        "not set": "NOT_SET",
    }
    if cleaned_lower in ROLE_MAP:
        return ROLE_MAP[cleaned_lower], None
    return "NOT_SET", f"Client role '{s}' unrecognized; set to NOT_SET."


def generate_import_preview(
    tenant_id: str,
    raw_rows: list[list[Any]],
    mapping: dict[str, int],
    client_overrides: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Generate row-by-row preview with entity matching, date interpretations, and warnings."""
    today_d = timezone.now().astimezone(IST).date()
    overrides = client_overrides or {}

    # Pre-fetch existing active clients and matters for this tenant
    existing_clients = list(Client.objects.filter(tenant_id=tenant_id, status="ACTIVE"))
    existing_matters = list(Matter.objects.filter(tenant_id=tenant_id, status="ACTIVE"))
    existing_case_links = list(MatterCaseLink.objects.filter(tenant_id=tenant_id))

    # Pre-fetch active users for lead lawyer email matching
    active_users = {
        u.email.strip().lower(): u.user_id
        for u in AppUser.objects.filter(tenant_id=tenant_id, active=True)
        if u.email
    }

    # Index existing case links by matter_id
    # Index existing case links by matter_id
    case_links_by_matter: dict[str, list[MatterCaseLink]] = {}
    for cl in existing_case_links:
        case_links_by_matter.setdefault(cl.matter_id, []).append(cl)

    # Index matters by client_id and map client display names
    matters_by_client: dict[str, list[Matter]] = {}
    matter_client_map: dict[str, str] = {}
    for m in existing_matters:
        if m.client_id:
            matters_by_client.setdefault(m.client_id, []).append(m)
            matter_client_map[m.matter_id] = m.client_id
    client_display_map: dict[str, str] = {c.client_id: c.display_name for c in existing_clients}

    # Pre-fetch existing hearing dates by matter_id
    hearings_by_matter: dict[str, set[date]] = defaultdict(set)
    for h in Hearing.objects.filter(tenant_id=tenant_id).values("matter_id", "date"):
        hearings_by_matter[h["matter_id"]].add(h["date"])

    # Build case link lookup dicts once (Session R3 Fix Pass 2 - Item 2)
    case_number_to_client: dict[str, dict[str, str]] = {}
    raw_case_to_client: dict[str, dict[str, str]] = {}
    for cl in existing_case_links:
        c_id = matter_client_map.get(cl.matter_id)
        if c_id and cl.value_normalized:
            c_name = client_display_map.get(c_id, c_id)
            case_number_to_client[cl.value_normalized] = {
                "client_id": c_id,
                "client_name": c_name,
                "matter_id": cl.matter_id,
            }
            raw_case_to_client[cl.value_normalized.strip().lower()] = {
                "client_id": c_id,
                "client_name": c_name,
                "matter_id": cl.matter_id,
            }

    preview_rows: list[dict[str, Any]] = []
    seen_in_batch_clients: dict[str, str] = {}  # norm_client_name -> virtual or existing client_id
    seen_in_batch_matters: set[tuple[str, str]] = set()  # (client_key, matter_key)

    total_new_clients = 0
    total_existing_clients = 0
    total_new_matters = 0
    total_existing_matters = 0
    total_hearings = 0
    total_warnings = 0

    def get_row_val(row_cells: list[Any], field: str) -> Any:
        col_idx = mapping.get(field)
        if col_idx is not None and 0 <= col_idx < len(row_cells):
            return row_cells[col_idx]
        return None

    for idx, row in enumerate(raw_rows):
        row_num = idx + 1
        row_warnings: list[str] = []

        client_name = str(get_row_val(row, "client_name") or "").strip()
        client_type = str(get_row_val(row, "client_type") or "COMPANY").strip().upper()
        if client_type not in ("INDIVIDUAL", "COMPANY", "LLP", "PARTNERSHIP", "GOVERNMENT", "OTHER"):
            client_type = "COMPANY"

        cin_pan_raw = str(get_row_val(row, "cin_or_pan") or "").strip()
        client_role_raw = get_row_val(row, "client_role")
        client_role, role_warning = normalize_client_role(client_role_raw)
        if role_warning:
            row_warnings.append(role_warning)

        matter_title = str(get_row_val(row, "matter_title") or "").strip()
        case_type = str(get_row_val(row, "case_type") or "").strip() or None
        court_or_forum = str(get_row_val(row, "court_or_forum") or "").strip() or None
        case_numbers_raw = str(get_row_val(row, "case_number") or "").strip()
        next_date_raw = get_row_val(row, "next_hearing_date")
        lead_lawyer_email = str(get_row_val(row, "lead_lawyer_email") or "").strip()
        notes = str(get_row_val(row, "notes") or "").strip() or None

        if not client_name:
            row_warnings.append("Missing client name; row cannot be imported.")
        if not matter_title:
            row_warnings.append("Missing matter title; row cannot be imported.")

        # Parse CIN or PAN
        cin = None
        pan = None
        if cin_pan_raw:
            cleaned_id = re.sub(r"[^A-Za-z0-9]", "", cin_pan_raw).upper()
            if len(cleaned_id) == 21:
                cin = cleaned_id
            elif len(cleaned_id) == 10:
                pan = cleaned_id
            else:
                row_warnings.append(f"Identifier '{cin_pan_raw}' is neither a 21-char CIN nor a 10-char PAN.")

        # --- 1. Client Matching ---
        norm_client_name = normalize_client_name(client_name)
        matched_client: Client | None = None
        override = (
            overrides.get(str(row_num))
            or overrides.get(str(idx))
        )
        if override and override.get("action") == "USE_EXISTING":
            matched_client = next(
                (c for c in existing_clients if c.client_id == override.get("client_id")),
                None,
            )
        elif not (override and override.get("action") == "CREATE_NEW"):
            # Step A: Match by CIN or PAN
            if cin:
                matched_client = next(
                    (c for c in existing_clients if (c.identifiers or {}).get("cin") == cin),
                    None,
                )
            if not matched_client and pan:
                matched_client = next(
                    (c for c in existing_clients if (c.identifiers or {}).get("pan") == pan),
                    None,
                )
            # Step B: Match by folded normalized display_name (case, punctuation, M/s, Pvt/Ltd)
            if not matched_client and norm_client_name:
                matched_client = next(
                    (c for c in existing_clients if normalize_client_name(c.display_name) == norm_client_name),
                    None,
                )

        client_status: str
        resolved_client_id: str | None
        if matched_client:
            client_status = "EXISTING_CLIENT"
            resolved_client_id = matched_client.client_id
            total_existing_clients += 1
        else:
            client_status = "NEW_CLIENT"
            resolved_client_id = seen_in_batch_clients.get(norm_client_name)
            if not resolved_client_id:
                resolved_client_id = f"batch_new_cli_{idx}"
                seen_in_batch_clients[norm_client_name] = resolved_client_id
                total_new_clients += 1

        # --- 2. Case Numbers Normalization ---
        parsed_case_numbers: list[dict[str, Any]] = []
        canonical_case_numbers: list[str] = []
        raw_case_numbers: list[str] = []

        if case_numbers_raw:
            parts = [p.strip() for p in case_numbers_raw.split(";") if p.strip()]
            for p in parts:
                norm = normalize_case_number(p)
                if norm.parseable:
                    parsed_case_numbers.append(
                        {
                            "raw": p,
                            "normalized": norm.canonical,
                            "parseable": True,
                        }
                    )
                    canonical_case_numbers.append(norm.canonical)
                else:
                    parsed_case_numbers.append(
                        {
                            "raw": p,
                            "normalized": p,
                            "parseable": False,
                        }
                    )
                    raw_case_numbers.append(p.strip().lower())
                    row_warnings.append(
                        f"Case number '{p}' could not be canonically parsed; it will not be tracked automatically."
                    )

        # Check if case number exists under a different client (Session R3 Fix Pass 2 - Item 2)
        matched_other_client_case: dict[str, str] | None = None
        for c_norm in canonical_case_numbers:
            found = case_number_to_client.get(c_norm)
            if found and found["client_id"] != resolved_client_id:
                matched_other_client_case = found
                row_warnings.append(
                    f"Case number '{c_norm}' already exists under a different client '{found['client_name']}'; not auto-matched."
                )
                break

        if not matched_other_client_case and raw_case_numbers:
            for r_norm in raw_case_numbers:
                found = raw_case_to_client.get(r_norm)
                if found and found["client_id"] != resolved_client_id:
                    matched_other_client_case = found
                    row_warnings.append(
                        f"Case number '{r_norm}' already exists under a different client '{found['client_name']}'; not auto-matched."
                    )
                    break

        # --- 3. Matter Deduplication & Matching ---
        # Rule 8:
        # - readable case number: the normalized case number
        # - unreadable case number: trimmed raw text, compared case-insensitively
        # - no case number: exact title -> flag as possible duplicate (matched by title)
        matched_matter: Matter | None = None
        match_reason: str | None = None

        if matched_other_client_case and client_status == "NEW_CLIENT":
            other_c_name = matched_other_client_case["client_name"]
            row_warnings.append(f"possible duplicate (case number exists under {other_c_name})")
            if not (override and override.get("action") in ("CREATE_ANYWAY", "IMPORT_ANYWAY", "CREATE_NEW")):
                matched_matter = Matter(
                    title=matter_title,
                    matter_id=matched_other_client_case["matter_id"],
                    client_id=matched_other_client_case["client_id"],
                )
                match_reason = f"case number exists under {other_c_name}"

        candidate_matters = matters_by_client.get(resolved_client_id or "", [])
        if not matched_matter and canonical_case_numbers:
            for m in candidate_matters:
                m_links = case_links_by_matter.get(m.matter_id, [])
                for cl in m_links:
                    if cl.value_normalized in canonical_case_numbers:
                        matched_matter = m
                        match_reason = f"case number {cl.value_normalized}"
                        break
                if matched_matter:
                    break

        if not matched_matter and raw_case_numbers:
            for m in candidate_matters:
                m_links = case_links_by_matter.get(m.matter_id, [])
                for cl in m_links:
                    if cl.value_normalized.strip().lower() in raw_case_numbers:
                        matched_matter = m
                        match_reason = f"raw case number '{cl.value_normalized}'"
                        break
                if matched_matter:
                    break

        if not matched_matter and not parsed_case_numbers and matter_title:
            for m in candidate_matters:
                if m.title.strip().lower() == matter_title.lower():
                    row_warnings.append(
                        f"Possible duplicate (matched by title with matter '{m.matter_id}')"
                    )
                    if not (override and override.get("action") in ("CREATE_ANYWAY", "IMPORT_ANYWAY", "CREATE_NEW")):
                        matched_matter = m
                        match_reason = "exact title"
                    break

        # Check duplicate within batch
        batch_key = (
            resolved_client_id or norm_client_name,
            canonical_case_numbers[0]
            if canonical_case_numbers
            else (raw_case_numbers[0] if raw_case_numbers else matter_title.lower()),
        )
        if batch_key in seen_in_batch_matters:
            matched_matter = matched_matter or Matter(title=matter_title)
            match_reason = match_reason or "duplicate row within uploaded spreadsheet"
            row_warnings.append("Duplicate row within uploaded spreadsheet; will skip.")
        else:
            seen_in_batch_matters.add(batch_key)

        # --- 4. Date Parsing ---
        parsed_date, date_warnings = parse_import_date(next_date_raw, today_d)
        row_warnings.extend(date_warnings)

        matter_status: str
        if matched_matter:
            matter_status = "ALREADY_EXISTS"
            total_existing_matters += 1
            # Warning: next hearing date was not added if matter already exists (Session R3 Fix Pass - Item 17)
            if parsed_date:
                m_hearings = hearings_by_matter.get(getattr(matched_matter, "matter_id", ""), set())
                if parsed_date not in m_hearings:
                    row_warnings.append(
                        f"matter already exists; next date {parsed_date.isoformat()} was not added"
                    )
        else:
            matter_status = "NEW_MATTER"
            total_new_matters += 1
            if parsed_date:
                total_hearings += 1

        # --- 5. Lead Lawyer Email Matching ---
        lead_user_id: str | None = None
        if lead_lawyer_email:
            lead_user_id = active_users.get(lead_lawyer_email.lower())
            if not lead_user_id:
                row_warnings.append(
                    f"Lead lawyer email '{lead_lawyer_email}' does not match any active user in this firm; importer will be sole LEAD."
                )

        total_warnings += len(row_warnings)

        preview_rows.append(
            {
                "row_number": row_num,
                "client_name": client_name,
                "client_type": client_type,
                "client_role": client_role,
                "client_status": client_status,
                "matched_client_id": matched_client.client_id if matched_client else None,
                "cin": cin,
                "pan": pan,
                "matter_title": matter_title,
                "matter_status": matter_status,
                "matched_matter_id": getattr(matched_matter, "matter_id", None) if matched_matter else None,
                "match_reason": match_reason,
                "case_type": case_type,
                "court_or_forum": court_or_forum,
                "case_numbers": parsed_case_numbers,
                "hearing_date": parsed_date.isoformat() if parsed_date else None,
                "lead_lawyer_email": lead_lawyer_email,
                "lead_user_id": lead_user_id,
                "notes": notes,
                "warnings": row_warnings,
            }
        )

    return {
        "summary": {
            "total_rows": len(raw_rows),
            "new_clients": total_new_clients,
            "existing_clients": total_existing_clients,
            "new_matters": total_new_matters,
            "existing_matters": total_existing_matters,
            "hearings_to_create": total_hearings,
            "warnings_count": total_warnings,
        },
        "rows": preview_rows,
    }


def execute_import_commit(
    tenant_id: str,
    importer_user_id: str,
    import_id: str,
    preview_data: dict[str, Any] | None = None,
    original_filename: str | None = None,
) -> dict[str, Any]:
    """Execute import creation in a single transaction with advisory lock per tenant."""
    now = timezone.now()
    counts = {
        "clients_created": 0,
        "matters_created": 0,
        "hearings_created": 0,
        "case_links_created": 0,
        "skipped_existing_matters": 0,
    }
    report: list[dict[str, Any]] = []

    # Cache created clients within this transaction (normalized name -> client_id)
    created_clients_cache: dict[str, str] = {}

    with transaction.atomic():
        with connection.cursor() as cursor:
            # Set local transaction tenant and user context for RLS
            cursor.execute("SELECT set_config('app.tenant_id', %s, true)", [tenant_id])
            cursor.execute("SELECT set_config('app.user_id', %s, true)", [importer_user_id])

            # Acquire per-tenant advisory lock
            cursor.execute(
                "SELECT pg_advisory_xact_lock(hashtext(%s))",
                [f"import_{tenant_id}"],
            )

            # Check that job is not already COMMITTED with FOR UPDATE
            cursor.execute(
                "SELECT status, file_ref, filename, mapping, client_overrides, sheet_name FROM tpl.import_job WHERE tenant_id = %s AND import_id = %s FOR UPDATE",
                [tenant_id, import_id],
            )
            job_row = cursor.fetchone()
            if not job_row:
                raise ValueError("Import job not found.")
            if job_row[0] == "COMMITTED":
                raise ValueError("This import job has already been committed.")

            job_file_ref = job_row[1]
            job_filename = job_row[2]
            job_mapping = job_row[3] or {}
            job_client_overrides = job_row[4] or {}
            job_sheet_name = job_row[5]

        actual_filename = original_filename or job_filename

        # Re-parse file and re-evaluate duplicate check INSIDE advisory lock (Session R3 Fix Pass - Item 2 & 3)
        if preview_data is not None:
            actual_preview = preview_data
        else:
            from parse.vault import VaultStorage
            storage = VaultStorage()
            file_bytes = storage.get_raw(job_file_ref)
            file_type, sheet_names, _ = inspect_file_safety(file_bytes, job_filename)
            selected_sheet = job_sheet_name or (sheet_names[0] if sheet_names else None)
            _, raw_rows, _ = parse_raw_rows(file_bytes, file_type, sheet_name=selected_sheet)
            actual_preview = generate_import_preview(
                tenant_id=tenant_id,
                raw_rows=raw_rows,
                mapping=job_mapping,
                client_overrides=job_client_overrides,
            )

        for row in actual_preview.get("rows", []):
            row_num = row.get("row_number", 0)
            client_name = row.get("client_name")
            matter_title = row.get("matter_title")
            warnings = row.get("warnings", [])

            if not client_name or not matter_title:
                report.append(
                    {
                        "row": row_num,
                        "status": "SKIPPED",
                        "reason": "Missing required client or matter title",
                        "warnings": warnings,
                    }
                )
                continue

            # Skip if matter already exists
            if row.get("matter_status") == "ALREADY_EXISTS":
                counts["skipped_existing_matters"] += 1
                report.append(
                    {
                        "row": row_num,
                        "status": "SKIPPED_DUPLICATE",
                        "client_id": row.get("matched_client_id"),
                        "matter_id": row.get("matched_matter_id"),
                        "reason": row.get("match_reason") or "Matter already exists",
                        "warnings": warnings,
                    }
                )
                continue

            # 1. Resolve or Create Client
            norm_name = normalize_client_name(client_name)
            client_id = row.get("matched_client_id")
            if not client_id:
                client_id = created_clients_cache.get(norm_name)

            if not client_id:
                identifiers = {}
                if row.get("cin"):
                    identifiers["cin"] = row.get("cin")
                if row.get("pan"):
                    identifiers["pan"] = row.get("pan")

                new_client = create_client_service(
                    tenant_id=tenant_id,
                    actor_id=importer_user_id,
                    display_name=client_name,
                    client_type=row.get("client_type", "COMPANY"),
                    identifiers=identifiers,
                    contact={},
                    notes=f"Imported from {actual_filename}, row {row_num}",
                    import_id=import_id,
                )
                client_id = new_client["client_id"]
                created_clients_cache[norm_name] = client_id
                counts["clients_created"] += 1

            # 2. Create Matter
            new_matter = create_matter_service(
                tenant_id=tenant_id,
                actor_id=importer_user_id,
                client_id=client_id,
                title=matter_title,
                client_role=row.get("client_role") or "NOT_SET",
                case_type=row.get("case_type"),
                court_or_forum=row.get("court_or_forum"),
                jurisdiction_state=None,
                client_matter_no=None,
                walled=False,
                additional_lead_user_id=row.get("lead_user_id"),
                import_id=import_id,
            )
            matter_id = new_matter["matter_id"]
            counts["matters_created"] += 1

            # 3. Create Case Links
            for c_info in row.get("case_numbers", []):
                raw_c = c_info.get("raw")
                if raw_c:
                    create_case_link_service(
                        tenant_id=tenant_id,
                        actor_id=importer_user_id,
                        matter_id=matter_id,
                        value_raw=raw_c,
                        scheme="IN_COURT",
                        role="MAIN",
                        track=c_info.get("parseable", True),
                        import_id=import_id,
                    )
                    counts["case_links_created"] += 1

            # 4. Create Hearing if date present
            h_date_str = row.get("hearing_date")
            hearing_id = None
            if h_date_str:
                h_date = date.fromisoformat(h_date_str)
                # Session R3 Fix Pass - Item 4: Leave outcome empty, provenance in source_ref
                new_hearing = create_hearing_service(
                    tenant_id=tenant_id,
                    actor_id=importer_user_id,
                    matter_id=matter_id,
                    hearing_date=h_date,
                    court_no=None,
                    item_no=None,
                    bench=row.get("court_or_forum"),
                    purpose="Listing",
                    source="MANUAL",
                    source_ref=f"import:{import_id}:row {row_num}",
                    status="LISTED",  # Correction 7: LISTED, never HEARD
                    outcome=None,
                    notes=None,
                    import_id=import_id,
                )
                hearing_id = new_hearing["hearing_id"]
                counts["hearings_created"] += 1

            report.append(
                {
                    "row": row_num,
                    "status": "CREATED",
                    "client_id": client_id,
                    "matter_id": matter_id,
                    "hearing_id": hearing_id,
                    "warnings": warnings,
                }
            )

        # Update import_job to COMMITTED
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE tpl.import_job
                SET status = 'COMMITTED', counts = %s, report = %s, committed_at = %s, committed_by = %s
                WHERE tenant_id = %s AND import_id = %s
                """,
                [
                    json.dumps(counts),
                    json.dumps(report),
                    now,
                    importer_user_id,
                    tenant_id,
                    import_id,
                ],
            )

        # Record IMPORT_COMMIT audit event
        record_audit_event(
            tenant_id=tenant_id,
            actor=importer_user_id,
            action="IMPORT_COMMIT",
            object_ref=import_id,
            decision="ALLOW",
            detail={
                "import_id": import_id,
                "filename": actual_filename,
                "counts": counts,
            },
        )

    return {
        "import_id": import_id,
        "status": "COMMITTED",
        "counts": counts,
        "report": report,
    }


def generate_import_report_csv(report: list[dict[str, Any]]) -> str:
    """Generate CSV report with formula sanitization."""
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow(["Row", "Status", "Client ID", "Matter ID", "Hearing ID", "Reason / Warnings"])
    for item in report:
        warnings_str = "; ".join(item.get("warnings", []))
        reason = item.get("reason", "")
        combined_notes = f"{reason} | {warnings_str}" if (reason and warnings_str) else (reason or warnings_str)

        writer.writerow(
            [
                sanitize_csv_cell(item.get("row")),
                sanitize_csv_cell(item.get("status")),
                sanitize_csv_cell(item.get("client_id")),
                sanitize_csv_cell(item.get("matter_id")),
                sanitize_csv_cell(item.get("hearing_id")),
                sanitize_csv_cell(combined_notes),
            ]
        )

    return output.getvalue()


def generate_sample_template_bytes(fmt: str = "xlsx") -> bytes:
    """Generate downloadable empty import template with headers."""
    if fmt == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(TEMPLATE_COLUMNS)
        return out.getvalue().encode("utf-8")

    # Excel template
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Live Cases"
    ws.append(TEMPLATE_COLUMNS)

    # Style header row slightly
    for col_num in range(1, len(TEMPLATE_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = openpyxl.styles.Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
